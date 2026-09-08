#!/usr/bin/env python3
"""Normaliza los tabulados regionales ENUSC 2025 para la historia web."""

import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


SOURCE = "https://www.ine.gob.cl/docs/default-source/seguridad-ciudadana/cuadros-estadisticos/2025/tabulados-regionales---enusc-2025.xlsx"


def theme(variable: str) -> str:
    if variable.startswith(("PAD", "P_FUENTE", "PCOS", "P_INSEG", "PED", "P_EXPOS")):
        return "Percepción y temor"
    if variable.startswith(("P_DESORDENES", "P_INCIVILIDADES", "PRESENCIA_TRAFICO", "PRESENCIA_ARMAS")):
        return "Entorno barrial"
    if variable.startswith("P_MOD_ACTIVIDADES"):
        return "Cambios de comportamiento"
    if variable.startswith(("EV_", "EVAL_", "PRESENCIA_CARABINEROS")):
        return "Instituciones y policías"
    if variable.startswith(("MEDIDAS_", "VECINOS_MEDIDAS_")):
        return "Protección y organización"
    if variable.startswith(("DEN_", "COSC_")):
        return "Denuncia y cifra oculta"
    return "Victimización"


def number(value):
    return round(float(value), 8) if isinstance(value, (int, float)) else None


def main():
    source_path = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/enusc-2025.xlsx")
    output_path = Path(sys.argv[2] if len(sys.argv) > 2 else "public/enusc-data.json")
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    index = workbook["Índice"]
    metadata = []
    for row in index.iter_rows(min_row=9, values_only=True):
        if not isinstance(row[0], int):
            continue
        variable = str(row[1])
        metadata.append({
            "order": row[0],
            "variable": variable,
            "title": str(row[2]).strip(),
            "type": row[3],
            "level": row[4],
            "disaggregation": row[5],
            "weight": row[6],
            "filter": row[7],
            "sample": row[8],
            "quality": {
                "national": number(row[9]),
                "nationalDisaggregated": number(row[10]),
                "regional": number(row[11]),
                "regionalDisaggregated": number(row[12]),
            },
            "theme": theme(variable),
        })

    tabulations = {}
    for item in metadata:
        ws = workbook[item["variable"]]
        rows = list(ws.iter_rows(values_only=True))
        headers = list(rows[3])
        category_column = len(headers) > 1 and headers[1] == "Categoría"
        start = 2 if category_column else 1
        groups = []
        for column in range(start, len(headers), 4):
            label = headers[column]
            if label is None:
                continue
            groups.append({"label": str(label), "column": column})
        records = []
        for row in rows[4:]:
            if not row or row[0] is None:
                continue
            region = str(row[0]).strip()
            category = str(row[1]).strip() if category_column and row[1] is not None else None
            estimates = []
            for group in groups:
                column = group["column"]
                estimate = number(row[column] if column < len(row) else None)
                if estimate is None:
                    continue
                raw_note = row[column + 3] if column + 3 < len(row) else None
                note = str(raw_note).strip() if raw_note not in (None, "") else None
                estimates.append({
                    "group": group["label"],
                    "estimate": estimate,
                    "lower": number(row[column + 1] if column + 1 < len(row) else None),
                    "upper": number(row[column + 2] if column + 2 < len(row) else None),
                    "note": note,
                })
            if estimates:
                records.append({"region": region, "category": category, "estimates": estimates})
        tabulations[item["variable"]] = records

    payload = {
        "year": 2025,
        "source": SOURCE,
        "metadata": metadata,
        "themes": dict(Counter(item["theme"] for item in metadata)),
        "tabulations": tabulations,
        "qualityNotes": {
            "1": "Estimación poco fiable (coeficiente de variación mayor a 15% y menor o igual a 30%. En el caso de estimaciones de razón, si no cumple con el umbral de aceptación asociado a su error estándar). Se recomienda utilizar con precaución esta estimación, ya que podría llevar a conclusiones poco acertadas.",
            "2": "Estimación no fiable (número de casos muestrales menor a 60, grados de libertad menores a 9 o coeficiente de variación mayor a 30%). No se recomienda el uso de esta estimación.",
        },
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({"variables": len(metadata), "tabulations": len(tabulations), "bytes": output_path.stat().st_size, "themes": payload["themes"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
