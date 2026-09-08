"""Convierte la hoja Mortalidad del libro oficial de estadísticas vitales a JSON."""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


# La fuente puede reemplazarse por una edición futura sin modificar el script.
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("series-vitales.xlsx")
OUTPUT = Path(__file__).resolve().parents[1] / "public" / "mortality-data.json"
worksheet = load_workbook(SOURCE, read_only=True, data_only=True)["Mortalidad"]
series = []

for row in worksheet.iter_rows(min_row=2, values_only=True):
    match = re.fullmatch(r"\s*(\d{4})(?:\(p\))?\s*", str(row[0]), re.IGNORECASE)
    if not match:
        continue
    series.append(
        {
            "year": int(match.group(1)),
            "provisional": "(p)" in str(row[0]).lower(),
            "crude": float(row[1]),
            "infant": float(row[2]),
            "neonatal": None if row[3] is None else float(row[3]),
            "fetal": None if row[4] is None else float(row[4]),
            "under5": float(row[5]),
            "lifeBoth": float(row[6]),
            "lifeMen": float(row[7]),
            "lifeWomen": float(row[8]),
        }
    )

OUTPUT.write_text(
    json.dumps({"source": SOURCE.name, "series": series}, ensure_ascii=False, separators=(",", ":")),
    encoding="utf-8",
)
print(f"{len(series)} años escritos en {OUTPUT}")
