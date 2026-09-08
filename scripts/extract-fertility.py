"""Convierte la hoja Fecundidad del libro oficial de estadísticas vitales a JSON."""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


# La ruta del libro puede entregarse por argumento para mantener el proceso reproducible.
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("series-vitales.xlsx")
OUTPUT = Path(__file__).resolve().parents[1] / "public" / "fertility-data.json"

worksheet = load_workbook(SOURCE, read_only=True, data_only=True)["Fecundidad"]
series = []

# Extrae únicamente las filas anuales; las notas al pie de la planilla quedan excluidas.
for row in worksheet.iter_rows(min_row=2, values_only=True):
    match = re.search(r"\d{4}", str(row[0])) if row[0] is not None else None
    if not match or row[1] is None or row[27] is None:
        continue
    series.append(
        {
            "year": int(match.group()),
            "provisional": "(p)" in str(row[0]).lower(),
            "population": int(row[1]),
            "correctedBirths": int(row[2]),
            "birthRate": float(row[3]),
            "women1549": int(row[4]),
            "generalRate": float(row[5]),
            "specificRates": [
                {"label": label, "value": float(row[index])}
                for label, index in zip(
                    ["15 a 19", "20 a 24", "25 a 29", "30 a 34", "35 a 39", "40 a 44", "45 a 49"],
                    range(20, 27),
                )
            ],
            "tgf": float(row[27]),
            "tbr": float(row[28]),
            "meanFertilityAge": float(row[29]),
            "meanMotherAge": float(row[30]),
            "medianMotherAge": float(row[31]),
        }
    )

OUTPUT.write_text(
    json.dumps({"source": SOURCE.name, "series": series}, ensure_ascii=False, separators=(",", ":")),
    encoding="utf-8",
)
print(f"{len(series)} años escritos en {OUTPUT}")
