"""Convierte las hojas nacionales AS, H y M en series optimizadas para el gráfico web."""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


INDICATORS = [
    ("pet", "Población en edad de trabajar (Total)", 3, 4, "level"),
    ("labor", "Fuerza de trabajo (Total)", 5, 6, "level"),
    ("employed", "Población ocupada (Total)", 7, 8, "level"),
    ("unemployed", "Población desocupada (Total)", 9, 10, "level"),
    ("ceased", "Población desocupada (Cesantes)", 11, 12, "level"),
    ("firstJob", "Población desocupada (Buscan trabajo por primera vez)", 13, 14, "level"),
    ("inactive", "Fuera de la fuerza de trabajo (Total)", 15, 16, "level"),
    ("initiators", "Fuera de la fuerza de trabajo (Iniciadores)", 17, 18, "level"),
    ("potential", "Fuera de la fuerza de trabajo (Inactivos potencialmente activos)", 19, 20, "level"),
    ("habitual", "Fuera de la fuerza de trabajo (Inactivos habituales)", 21, 22, "level"),
    ("unemploymentRate", "Tasa de desocupación [1]", 23, 24, "rate"),
    ("employmentRate", "Tasa de ocupación [2]", 25, 26, "rate"),
    ("participation", "Tasa de participación [3]", 27, 28, "rate"),
]


def clean_note(value):
    """Conserva únicamente las marcas estadísticas a y b."""
    match = re.search(r"[ab]", str(value or "").strip().lower())
    return match.group(0) if match else None


def main(source: str, destination: str):
    """Lee las tres hojas y genera un JSON consistente por período e indicador."""
    workbook = load_workbook(source, data_only=True, read_only=True)
    sexes = {"AS": "Total", "H": "Hombres", "M": "Mujeres"}
    output = {
        "indicators": [{"id": i, "label": label, "unit": unit} for i, label, _, _, unit in INDICATORS],
        "series": {},
    }
    for sheet_name, sex_label in sexes.items():
        sheet = workbook[sheet_name]
        points = []
        for row in sheet.iter_rows(min_row=8, values_only=True):
            if not isinstance(row[0], (int, float)) or not row[1]:
                continue
            point = {"year": int(row[0]), "quarter": str(row[1]).strip(), "values": {}}
            for indicator_id, _, note_col, value_col, _ in INDICATORS:
                value = row[value_col - 1]
                point["values"][indicator_id] = {
                    "value": float(value) if isinstance(value, (int, float)) else None,
                    "note": clean_note(row[note_col - 1]),
                }
            points.append(point)
        output["series"][sex_label] = points
    Path(destination).write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
