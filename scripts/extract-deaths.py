"""Convierte las hojas Defunciones y Nacimientos del libro oficial a JSON analítico."""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


# La ruta se recibe por argumento para poder regenerar el archivo con futuras ediciones del libro.
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("series-vitales.xlsx")
OUTPUT = Path(__file__).resolve().parents[1] / "public" / "deaths-data.json"
workbook = load_workbook(SOURCE, read_only=True, data_only=True)


def annual_rows(sheet_name):
    """Conserva solamente filas cuyo primer campo identifica inequívocamente un año."""
    rows = {}
    for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True):
        match = re.fullmatch(r"\s*(\d{4})(?:\(p\))?\s*", str(row[0]), re.IGNORECASE)
        if match:
            rows[int(match.group(1))] = row
    return rows


deaths = annual_rows("Defunciones")
births = annual_rows("Nacimientos")
series = []

for year, row in deaths.items():
    birth_row = births[year]
    observed_births = int(birth_row[1])
    known_age_births = sum(int(birth_row[index] or 0) for index in range(7, 17))
    young_births = int(birth_row[7] or 0) + int(birth_row[8] or 0)

    # Las razones usan nacimientos observados porque la serie corregida no está publicada para 2023-2024.
    def per_thousand(value):
        return None if value is None else float(value) / observed_births * 1000

    series.append(
        {
            "year": year,
            "provisional": "(p)" in str(row[0]).lower(),
            "total": int(row[1]),
            "men": int(row[2]),
            "women": int(row[3]),
            # La serie de sexo indeterminado se excluye deliberadamente por solicitud editorial.
            "masculinity": float(row[5]),
            "neonatal": None if row[6] is None else int(row[6]),
            "infant": None if row[7] is None else int(row[7]),
            "age1to4": None if row[8] is None else int(row[8]),
            "fetal": None if row[9] is None else int(row[9]),
            "observedBirths": observed_births,
            "youngMotherShare": young_births / known_age_births * 100,
            "neonatalPerThousandBirths": per_thousand(row[6]),
            "infantPerThousandBirths": per_thousand(row[7]),
            "age1to4PerThousandBirths": per_thousand(row[8]),
            "fetalPerThousandBirths": per_thousand(row[9]),
        }
    )

OUTPUT.write_text(
    json.dumps({"source": SOURCE.name, "series": series}, ensure_ascii=False, separators=(",", ":")),
    encoding="utf-8",
)
print(f"{len(series)} años escritos en {OUTPUT}")
