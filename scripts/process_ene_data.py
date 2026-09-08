"""Convierte los libros oficiales de la ENE en un JSON liviano para la web."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


UPLOAD = Path("/workspace/scratch/60166223c83f/upload")
OUTPUT = Path(__file__).resolve().parents[1] / "public" / "ene-data.json"


def clean_number(value: Any) -> float | None:
    """Devuelve números redondeados y descarta códigos de calidad no numéricos."""
    return round(float(value), 4) if isinstance(value, (int, float)) else None


def principal_series() -> dict[str, list[dict[str, Any]]]:
    """Extrae las series nacionales para ambos sexos, hombres y mujeres."""
    workbook = load_workbook(UPLOAD / "indicadores_principales (1).xlsx", read_only=True, data_only=True)
    result: dict[str, list[dict[str, Any]]] = {}
    for sheet, label in (("AS", "Total"), ("M", "Mujeres"), ("H", "Hombres")):
        ws = workbook[sheet]
        rows: list[dict[str, Any]] = []
        for row in ws.iter_rows(min_row=8, values_only=True):
            year, quarter = row[0], row[1]
            if not isinstance(year, int) or not isinstance(quarter, str):
                continue
            pet, labor, employed, unemployed = map(clean_number, (row[3], row[5], row[7], row[9]))
            ceased, first_job = map(clean_number, (row[11], row[13]))
            unemployment_rate, employment_rate, participation = map(clean_number, (row[23], row[25], row[27]))
            if None in (pet, labor, employed, unemployed, unemployment_rate, employment_rate, participation):
                continue
            rows.append({
                "year": year,
                "quarter": " ".join(quarter.split()),
                "pet": pet,
                "labor": labor,
                "employed": employed,
                "unemployed": unemployed,
                "ceased": ceased,
                "firstJob": first_job,
                # Se conservan las tasas oficiales del libro (cuatro decimales)
                # para evitar un doble redondeo antes de mostrarlas a una cifra.
                "participation": participation,
                "employmentRate": employment_rate,
                "unemploymentRate": unemployment_rate,
            })
        result[label] = rows
    return result


def latest_breakdown(filename: str, sheet: str) -> dict[str, Any]:
    """Extrae la última distribución disponible de ramas o categorías."""
    workbook = load_workbook(UPLOAD / filename, read_only=True, data_only=True)
    ws = workbook[sheet]
    header_row = 4 if filename in {"rama.xlsx", "categoria.xlsx"} else 6
    headers = [ws.cell(header_row, col).value for col in range(1, ws.max_column + 1)]
    latest = None
    for row in ws.iter_rows(min_row=header_row + 2, values_only=True):
        if isinstance(row[0], int) and isinstance(row[1], str):
            latest = row
    if latest is None:
        return {"period": "", "items": []}
    items = []
    # Los libros alternan columnas de estimación y nota; conservamos las estimaciones numéricas.
    for index in range(2, len(latest)):
        value = clean_number(latest[index])
        header = headers[index] if index < len(headers) else None
        if value is not None and isinstance(header, str) and header.strip():
            items.append({"label": " ".join(header.split()), "value": value})
    return {"period": f"{latest[1]} {latest[0]}", "items": items[:18]}


def seasonal_unemployment() -> list[dict[str, Any]]:
    """Obtiene la tasa desestacionalizada usando la última versión disponible del ajuste."""
    workbook = load_workbook(UPLOAD / "ajuste_estacional_historico.xlsx", read_only=True, data_only=True)
    ws = workbook["tasa_as"]
    records = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        year, quarter = row[0], row[1]
        adjusted = clean_number(row[155]) if len(row) >= 156 else None
        if isinstance(year, int) and isinstance(quarter, str) and adjusted is not None:
            records.append({"year": year, "quarter": " ".join(quarter.split()), "value": adjusted})
    # El libro está en orden descendente; el sitio usa orden cronológico.
    return list(reversed(records))


def sector_contributions() -> list[dict[str, Any]]:
    """Calcula los tres sectores con mayor incidencia positiva anual en la ocupación."""
    workbook = load_workbook(UPLOAD / "rama.xlsx", read_only=True, data_only=True)
    ws = workbook["AS"]
    names = []
    for col in range(5, ws.max_column + 1, 2):
        header = ws.cell(6, col).value
        if isinstance(header, str):
            names.append((col + 1, " ".join(header.split())))
    rows: dict[tuple[int, str], tuple[Any, ...]] = {}
    for row in ws.iter_rows(min_row=8, values_only=True):
        if isinstance(row[0], int) and isinstance(row[1], str):
            rows[(row[0], " ".join(row[1].split()))] = row
    output = []
    for (year, quarter), row in rows.items():
        previous = rows.get((year - 1, quarter))
        if previous is None or not isinstance(previous[3], (int, float)):
            continue
        sectors = []
        for excel_col, label in names:
            index = excel_col - 1
            current_value = row[index] if index < len(row) else None
            previous_value = previous[index] if index < len(previous) else None
            if not isinstance(current_value, (int, float)) or not isinstance(previous_value, (int, float)):
                continue
            change = (current_value / previous_value - 1) * 100 if previous_value else 0
            incidence = (current_value - previous_value) / previous[3] * 100
            if incidence > 0:
                sectors.append({"label": label, "change": round(change, 2), "incidence": round(incidence, 3)})
        sectors.sort(key=lambda item: item["incidence"], reverse=True)
        output.append({"year": year, "quarter": quarter, "items": sectors[:3]})
    return output


def category_contributions() -> list[dict[str, Any]]:
    """Calcula las categorías ocupacionales específicas con mayor incidencia positiva anual."""
    workbook = load_workbook(UPLOAD / "categoria.xlsx", read_only=True, data_only=True)
    ws = workbook["AS"]
    excluded = {
        "Independientes (Total) [2]",
        "Dependientes (Total) [3]",
        "Asalariados/as (Total) [4]",
    }
    display_names = {
        "Independientes (Empleadores/as)": "Personas empleadoras",
        "Independientes (Trabajadores/as por cuenta propia)": "Trabajadores/as por cuenta propia",
        "Independientes (Familiares no remunerados)": "Familiares no remunerados",
        "Asalariados/as (Sector privado)": "Personas asalariadas del sector privado",
        "Asalariados/as (Sector público) [5]": "Personas asalariadas del sector público",
        "Personal de servicio doméstico (Total) [6]": "Personal de servicio doméstico",
    }
    names = []
    for col in range(5, ws.max_column + 1, 2):
        header = ws.cell(6, col).value
        if isinstance(header, str) and " ".join(header.split()) not in excluded:
            normalized = " ".join(header.split())
            names.append((col + 1, display_names.get(normalized, normalized)))
    rows: dict[tuple[int, str], tuple[Any, ...]] = {}
    for row in ws.iter_rows(min_row=8, values_only=True):
        if isinstance(row[0], int) and isinstance(row[1], str):
            rows[(row[0], " ".join(row[1].split()))] = row
    output = []
    for (year, quarter), row in rows.items():
        previous = rows.get((year - 1, quarter))
        if previous is None or not isinstance(previous[3], (int, float)):
            continue
        items = []
        for excel_col, label in names:
            index = excel_col - 1
            current_value = row[index] if index < len(row) else None
            previous_value = previous[index] if index < len(previous) else None
            if not isinstance(current_value, (int, float)) or not isinstance(previous_value, (int, float)) or not previous_value:
                continue
            incidence = (current_value - previous_value) / previous[3] * 100
            if incidence > 0:
                items.append({
                    "label": label,
                    "change": round((current_value / previous_value - 1) * 100, 2),
                    "incidence": round(incidence, 3),
                })
        items.sort(key=lambda item: item["incidence"], reverse=True)
        output.append({"year": year, "quarter": quarter, "items": items[:3]})
    return output


def main() -> None:
    """Construye el archivo de datos consumido por la interfaz."""
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    sectors = sector_contributions()
    # El libro de ramas adjunto termina en febrero-abril; se completa el último período con el boletín 332.
    sectors.append({"year": 2026, "quarter": "Mar - May", "items": [
        {"label": "Industrias manufactureras", "change": 5.2, "incidence": None},
        {"label": "Actividades de atención de la salud humana y de asistencia social", "change": 6.0, "incidence": None},
        {"label": "Actividades profesionales, científicas y técnicas", "change": 11.6, "incidence": None},
    ]})
    categories = category_contributions()
    categories.append({"year": 2026, "quarter": "Mar - May", "items": [
        {"label": "Trabajadores/as por cuenta propia", "change": 5.3, "incidence": None},
        {"label": "Personas asalariadas informales", "change": 7.1, "incidence": None},
        {"label": "Personas empleadoras", "change": 3.8, "incidence": None},
    ]})
    payload = {
        "series": principal_series(),
        "seasonal": seasonal_unemployment(),
        "sectorContributions": sectors,
        "categoryContributions": categories,
        # Esta desagregación no está incluida en los Excel adjuntos. El único
        # registro disponible proviene del boletín ENE nacional N° 332.
        "absentEmployment": [{
            "year": 2026,
            "quarter": "Mar - May",
            "share": 4.7,
            "change": -6.0,
            "changePeople": -28083,
            "presentChange": 1.1,
        }],
        "branches": latest_breakdown("rama.xlsx", "AS"),
        "categories": latest_breakdown("categoria.xlsx", "AS"),
        "metadata": {
            "source": "Instituto Nacional de Estadísticas de Chile, Encuesta Nacional de Empleo",
            "updated": "30-06-2026",
            "qualityNotes": {
                "a": "Estimación poco fiable; debe utilizarse con precaución.",
                "b": "Estimación no fiable.",
            },
        },
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


if __name__ == "__main__":
    main()
