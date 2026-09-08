"""Extrae las divisiones manufactureras del cuadro oficial para la visualización web."""

import json
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path("/workspace/scratch/60166223c83f/library-downloads/industria-manufacturera-xlsx.xlsx")
OUTPUT = Path(__file__).resolve().parents[1] / "public" / "ippman-divisions.json"

workbook = load_workbook(SOURCE, read_only=True, data_only=True)
worksheet = workbook["IPP_Manufactura"]
rows = []

# El filtro reproduce el nivel de desagregación solicitado: división informada y grupo vacío.
for row in worksheet.iter_rows(min_row=6, values_only=True):
    year, month, division, group, _class, subclass, product, label, _, index, monthly, accumulated, annual = row[:13]
    if division is None or group is not None:
        continue
    rows.append({
        "year": int(year),
        "month": int(month),
        "division": int(division),
        "label": str(label),
        "index": float(index),
        "monthly": float(monthly),
        "accumulated": float(accumulated),
        "annual": float(annual),
    })

OUTPUT.write_text(json.dumps(rows, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(f"{len(rows)} observaciones escritas en {OUTPUT}")
