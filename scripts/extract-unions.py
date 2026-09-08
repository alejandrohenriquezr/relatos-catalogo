"""Integra las hojas Matrimonios y AUC del libro oficial en un JSON analítico."""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("series-vitales.xlsx")
OUTPUT = Path(__file__).resolve().parents[1] / "public" / "unions-data.json"
workbook = load_workbook(SOURCE, read_only=True, data_only=True)
AGE_LABELS = ["Menores de 15", "15-19", "20-24", "25-29", "30-34", "35-39", "40-44", "45-49", "50-54", "55-59", "60-64", "65-69", "70 y más"]
MIDPOINTS = [14, 17, 22, 27, 32, 37, 42, 47, 52, 57, 62, 67, 72]


def year_from(value):
    match = re.fullmatch(r"\s*(\d{4})(?:\(p\))?\s*", str(value), re.IGNORECASE)
    return int(match.group(1)) if match else None


def age_profile(values):
    """Resume grupos de edad; la media es aproximada y usa puntos medios convencionales."""
    if not any(value is not None for value in values):
        return None
    counts = [int(value or 0) for value in values]
    total = sum(counts)
    return {
        "groups": [{"label": label, "value": value} for label, value in zip(AGE_LABELS, counts)],
        "modalGroup": AGE_LABELS[max(range(len(counts)), key=counts.__getitem__)],
        "approxMean": sum(value * midpoint for value, midpoint in zip(counts, MIDPOINTS)) / total,
    }


marriages = []
for row in workbook["Matrimonios"].iter_rows(min_row=2, values_only=True):
    year = year_from(row[0])
    if year is None:
        continue
    marriages.append({
        "year": year,
        "provisional": "(p)" in str(row[0]).lower(),
        "total": int(row[1]),
        "rate": float(row[2]),
        "menAge": age_profile(row[3:16]),
        "womenAge": age_profile(row[16:29]),
    })

auc = []
for row in workbook["AUC"].iter_rows(min_row=2, values_only=True):
    year = year_from(row[0])
    if year is None:
        continue
    auc.append({
        "year": year,
        "provisional": "(p)" in str(row[0]).lower(),
        "total": int(row[1]),
        "rate": float(row[2]),
        "differentSex": int(row[3]),
        "sameSex": int(row[6]),
        "sameSexMen": int(row[7]),
        "sameSexWomen": int(row[8]),
    })

OUTPUT.write_text(json.dumps({"source": SOURCE.name, "marriages": marriages, "auc": auc}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(f"{len(marriages)} años de matrimonios y {len(auc)} años de AUC escritos en {OUTPUT}")
