"""Convierte la hoja Nacimientos del libro oficial de estadísticas vitales a JSON."""

import json
import re
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path("/workspace/scratch/60166223c83f/series-vitales.xlsx")
OUTPUT = Path(__file__).resolve().parents[1] / "public" / "births-data.json"

worksheet = load_workbook(SOURCE, read_only=True, data_only=True)["Nacimientos"]
headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
age_columns = [(index, str(label).replace("Nacimientos de mujeres de ", "").replace("Nacimientos de mujeres ", "")) for index, label in enumerate(headers) if label and str(label).startswith("Nacimientos de mujeres") and "no especificada" not in str(label)]
series = []

for row in worksheet.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        break
    year = int(re.search(r"\d{4}", str(row[0])).group())
    series.append({
        "year": year,
        "provisional": "(p)" in str(row[0]),
        "observed": int(row[1]),
        "men": int(row[3]),
        "women": int(row[4]),
        "masculinity": float(row[6]),
        "ages": [{"label": label, "value": int(row[index] or 0)} for index, label in age_columns],
    })

OUTPUT.write_text(json.dumps({"source": SOURCE.name, "series": series}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(f"{len(series)} años y {len(age_columns)} grupos de edad escritos en {OUTPUT}")
